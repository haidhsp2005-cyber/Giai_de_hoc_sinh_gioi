import os
import subprocess
import tempfile
from pathlib import Path
from typing import Optional
# Lazy import docx, pypdf inside functions to speed up Vercel cold starts

def parse_docx(file_path: Path) -> str:
    """Trích xuất toàn bộ văn bản và bảng biểu từ file .docx"""
    import docx
    doc = docx.Document(str(file_path))
    content_parts = []
    
    for element in doc.element.body:
        # Nếu là paragraph
        if element.tag.endswith('p'):
            for p in doc.paragraphs:
                if p._p == element and p.text.strip():
                    content_parts.append(p.text)
                    break
        # Nếu là table
        elif element.tag.endswith('tbl'):
            for t in doc.tables:
                if t._tbl == element:
                    table_md = []
                    for row in t.rows:
                        row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                        table_md.append("| " + " | ".join(row_text) + " |")
                    if table_md:
                        # Thêm markdown separator
                        cols_count = len(t.rows[0].cells)
                        sep = "| " + " | ".join(["---"] * cols_count) + " |"
                        if len(table_md) > 1:
                            table_md.insert(1, sep)
                        content_parts.append("\n" + "\n".join(table_md) + "\n")
                    break

    # Nếu duyệt theo body chưa đủ, fallback duyệt paragraphs và tables
    if not content_parts:
        for p in doc.paragraphs:
            if p.text.strip():
                content_parts.append(p.text)
        for t in doc.tables:
            table_md = []
            for row in t.rows:
                row_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                table_md.append("| " + " | ".join(row_text) + " |")
            if table_md:
                content_parts.append("\n" + "\n".join(table_md) + "\n")
                
    return "\n\n".join(content_parts).strip()


def parse_pdf(file_path: Path) -> str:
    """Trích xuất text từ file PDF"""
    from pypdf import PdfReader
    reader = PdfReader(str(file_path))
    pages_text = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text and text.strip():
            pages_text.append(f"--- [Trang {i+1}] ---\n" + text.strip())
    
    if not pages_text:
        return "[Thông báo: File PDF không chứa text có thể trích xuất trực tiếp (có thể là PDF scan dạng ảnh). Vui lòng chuyển thành ảnh hoặc dán nội dung đề bài.]"
        
    return "\n\n".join(pages_text).strip()


import base64
import httpx
from config import get_gemini_keys_from_env

def parse_image_gemini_vision(image_path: Path) -> Optional[str]:
    """Sử dụng Gemini Vision Multimodal để đọc ảnh đề bài chính xác cao (hỗ trợ cả Render Linux và Windows)"""
    keys = get_gemini_keys_from_env()
    if not keys:
        return None
    
    ext = image_path.suffix.lower()
    mime = "image/png" if ext == ".png" else "image/jpeg"
    if ext == ".webp":
        mime = "image/webp"

    try:
        b64_data = base64.b64encode(image_path.read_bytes()).decode("utf-8")
        prompt = (
            "Bạn là chuyên gia bóc tách đề thi Học sinh giỏi Tin học. "
            "Hãy đọc và chép lại nguyên văn, chính xác 100% toàn bộ nội dung đề bài trong ảnh trên, "
            "bao gồm: Tên bài, Đề bài/Mô tả, Dữ liệu vào (Input), Dữ liệu ra (Output), "
            "Giới hạn/Ràng buộc (Subtasks), và các bộ Ví dụ (Test mẫu). "
            "Giữ nguyên các công thức toán KaTeX và định dạng bảng biểu nếu có."
        )
        for key in keys[:3]:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-flash-latest:generateContent?key={key}"
            payload = {
                "contents": [
                    {
                        "parts": [
                            {"text": prompt},
                            {
                                "inlineData": {
                                    "mimeType": mime,
                                    "data": b64_data
                                }
                            }
                        ]
                    }
                ]
            }
            try:
                res = httpx.post(url, json=payload, timeout=30.0)
                if res.status_code == 200:
                    cand = res.json().get("candidates", [])
                    if cand:
                        text = cand[0].get("content", {}).get("parts", [{}])[0].get("text", "")
                        if text and text.strip():
                            return text.strip()
            except Exception:
                continue
    except Exception as e:
        print(f"[Gemini Vision OCR Error]: {e}")
    return None


def parse_image_windows_ocr(image_path: Path) -> str:
    """Sử dụng Windows.Media.Ocr qua PowerShell để nhận diện chữ trong ảnh (Chỉ dành cho Windows)"""
    if os.name != 'nt':
        return ""
    ps_script = f"""
    Add-Type -AssemblyName System.Drawing
    [Windows.Globalization.Language,Windows.Globalization,ContentType=WindowsRuntime] | Out-Null
    [Windows.Graphics.Imaging.BitmapDecoder,Windows.Graphics.Imaging,ContentType=WindowsRuntime] | Out-Null
    [Windows.Media.Ocr.OcrEngine,Windows.Media.Ocr,ContentType=WindowsRuntime] | Out-Null
    [Windows.Storage.StorageFile,Windows.Storage,ContentType=WindowsRuntime] | Out-Null

    async function Get-Text {{
        param([string]$Path)
        $file = [Windows.Storage.StorageFile]::GetFileFromPathAsync($Path).GetAwaiter().GetResult()
        $stream = $file.OpenAsync([Windows.Storage.FileAccessMode]::Read).GetAwaiter().GetResult()
        $decoder = [Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream).GetAwaiter().GetResult()
        $bitmap = $decoder.GetSoftwareBitmapAsync().GetAwaiter().GetResult()
        
        $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages()
        if (-not $engine) {{
            $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage([Windows.Globalization.Language]::new("en-US"))
        }}
        
        $result = $engine.RecognizeAsync($bitmap).GetAwaiter().GetResult()
        return $result.Text
    }}

    Get-Text -Path '{str(image_path.resolve()).replace("'", "''")}'
    """
    try:
        res = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps_script],
            capture_output=True,
            text=True,
            timeout=25
        )
        if res.returncode == 0 and res.stdout.strip():
            return res.stdout.strip()
        else:
            return f"[Nhận diện hình ảnh: Vui lòng kiểm tra lại ảnh hoặc chỉnh sửa văn bản đề bài nếu nhận diện chưa đầy đủ.\nChi tiết: {res.stderr.strip()}]"
    except Exception as e:
        return f"[Lỗi nhận diện OCR hình ảnh: {e}]"


def extract_problem_content(file_path: Path) -> dict:
    """Tự động phân loại và trích xuất nội dung đề bài từ file"""
    suffix = file_path.suffix.lower()
    
    if suffix == ".docx":
        text = parse_docx(file_path)
        file_type = "docx"
    elif suffix == ".pdf":
        text = parse_pdf(file_path)
        file_type = "pdf"
    elif suffix in [".png", ".jpg", ".jpeg", ".webp", ".bmp"]:
        # 1. Thử nhận diện bằng Gemini Multimodal Vision trước (độ chính xác cao nhất)
        text = parse_image_gemini_vision(file_path)
        # 2. Nếu không có key hoặc lỗi, fallback sang Windows OCR (nếu đang chạy trên Windows)
        if not text and os.name == 'nt':
            text = parse_image_windows_ocr(file_path)
        if not text:
            text = "[Không thể nhận diện ảnh đề bài. Vui lòng dán trực tiếp nội dung đề bài hoặc kiểm tra cấu hình API Key.]"
        file_type = "image"
    elif suffix in [".txt", ".md"]:
        text = file_path.read_text(encoding="utf-8", errors="replace")
        file_type = "text"
    else:
        text = f"Định dạng {suffix} chưa được hỗ trợ."
        file_type = "unsupported"
        
    return {
        "filename": file_path.name,
        "file_type": file_type,
        "content": text
    }
